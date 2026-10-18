"""
Apply all required fixes and updates to UEHIS_Service_Learning_Report_Aditya_Godar.docx.
Specifically:
1. Populate Certificate of Completion & Approval (Page ii) with official certification paragraphs and 4-party signature table.
2. Fix 5 occurrences of duplicate 'Prof. Prof.' prefix.
3. Fix List of Figures Figure 34 entry from 'Prof. Dr. Suvarna Patil' to 'Prof. Darshan Talhande'.
4. Fix Section 1.2 pronoun and capitalization mismatch ('Our mentor' -> 'Prof. Talhande', 'her' -> 'his').
5. Fix Section 4.1 Problem Statement typo ('(UEHIS). that ingests' -> '(UEHIS), which ingests').
6. Fix Section 9 missing designation ('and Faculty Dr. Aditi Saxena' -> 'and Faculty Coordinator Dr. Aditi Saxena') and en-dashes.
7. Insert Section 2 subheadings (2.1, 2.2, 2.3) in the document body matching the Table of Contents.
8. Update Section 10 wording ('Each image placeholder below represents...' -> 'Each photograph below documents an authentic milestone...').
"""

import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

DOCX_PATH = "UEHIS_Service_Learning_Report_Aditya_Godar.docx"

COLOR_NAVY = RGBColor(23, 43, 92)       # #172B5C
COLOR_DARK_TEXT = RGBColor(30, 41, 59)  # #1E293B

def set_cell_margins(cell, top=80, bottom=80, left=60, right=60):
    """Set cell padding in twips."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_borders(cell, top="none", bottom="none", left="none", right="none"):
    """Set borders on a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    borders_xml = f'<w:tcBorders {nsdecls("w")}>'
    for side, val in [("top", top), ("left", left), ("bottom", bottom), ("right", right)]:
        borders_xml += f'<w:{side} w:val="{val}"/>'
    borders_xml += '</w:tcBorders>'
    tcPr.append(parse_xml(borders_xml))

def apply_all_updates():
    print(f"Loading {DOCX_PATH}...")
    doc = docx.Document(DOCX_PATH)
    print(f"Loaded: {len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables.")

    # -------------------------------------------------------------
    # 1. POPULATE CERTIFICATE PAGE (Page ii)
    # -------------------------------------------------------------
    print("\n--- 1. Populating Certificate Page (Page ii) ---")
    cert_head_p = None
    cert_head_idx = None
    for idx, p in enumerate(doc.paragraphs):
        if "CERTIFICATE OF COMPLETION & APPROVAL" in p.text:
            cert_head_p = p
            cert_head_idx = idx
            break

    if cert_head_p is None:
        raise ValueError("Could not find CERTIFICATE OF COMPLETION & APPROVAL heading!")

    print(f"Found Certificate Heading at Paragraph {cert_head_idx}.")
    
    # Paragraph after heading: P10 is empty, P11 has page break
    p_empty = doc.paragraphs[cert_head_idx + 1]
    p_break = doc.paragraphs[cert_head_idx + 2]
    
    # Verify p_break has page break
    has_page_br = any(b.tag.endswith('br') for r in p_break.runs for b in r._element)
    print(f"P{cert_head_idx + 1} text: {repr(p_empty.text)}, P{cert_head_idx + 2} has page break: {has_page_br}")

    # Create the 3 certificate body paragraphs
    cert_p1 = doc.add_paragraph()
    cert_p1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    cert_p1.paragraph_format.space_before = Pt(0)
    cert_p1.paragraph_format.space_after = Pt(7)
    cert_p1.paragraph_format.line_spacing = 1.15
    r = cert_p1.add_run(
        "This is to certify that the Service Learning project report titled \"Urban Ecosystem Health Intelligence System (UEHIS)\" "
        "is a bona fide record of work carried out by Aditya Godar (PRN: 24070126014), a student of Semester V, B.Tech. in Artificial "
        "Intelligence and Machine Learning, Class AIML A1, at Symbiosis Institute of Technology (SIT), Pune, affiliated with "
        "Symbiosis International (Deemed University), during the Academic Year 2026–27."
    )
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)
    r.font.color.rgb = COLOR_DARK_TEXT

    cert_p2 = doc.add_paragraph()
    cert_p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    cert_p2.paragraph_format.space_before = Pt(0)
    cert_p2.paragraph_format.space_after = Pt(7)
    cert_p2.paragraph_format.line_spacing = 1.15
    r = cert_p2.add_run(
        "The project work was conducted in association with SEER Bharat Foundation, Pune (License No. 147561), under the "
        "institutional mentorship of Prof. Darshan Talhande (Director, SEER Bharat Foundation) and the academic supervision "
        "of the Department of Artificial Intelligence and Machine Learning, Symbiosis Institute of Technology, Pune."
    )
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)
    r.font.color.rgb = COLOR_DARK_TEXT

    cert_p3 = doc.add_paragraph()
    cert_p3.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    cert_p3.paragraph_format.space_before = Pt(0)
    cert_p3.paragraph_format.space_after = Pt(7)
    cert_p3.paragraph_format.line_spacing = 1.15
    r = cert_p3.add_run(
        "The candidate has successfully completed 120 hours of documented Service Learning work across 15 weeks "
        "(5th July 2026 to 16th October 2026), adhering to the prescribed academic regulations, ethical guidelines, and attendance norms "
        "of the Institute. The technical documentation, empirical findings, and personal contributions presented in this report represent "
        "the student's authentic intellectual effort in Computer Vision & Modeling."
    )
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)
    r.font.color.rgb = COLOR_DARK_TEXT

    cert_sp = doc.add_paragraph()
    cert_sp.paragraph_format.space_before = Pt(16)
    cert_sp.paragraph_format.space_after = Pt(0)

    # Create Signature Table (2 rows x 2 cols)
    sig_tbl = doc.add_table(rows=2, cols=2)
    sig_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_tbl.autofit = False
    sig_widths = [Inches(3.25), Inches(3.25)]
    for row in sig_tbl.rows:
        for c_idx, w in enumerate(sig_widths):
            cell = row.cells[c_idx]
            cell.width = w
            set_cell_borders(cell, top="none", bottom="none", left="none", right="none")
            set_cell_margins(cell, top=80, bottom=80, left=60, right=60)

    # Row 0, Cell 0: Student Signature Block
    c00 = sig_tbl.cell(0, 0)
    p00 = c00.paragraphs[0]
    p00.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p00.paragraph_format.line_spacing = 1.15
    p00.paragraph_format.space_before = Pt(0)
    p00.paragraph_format.space_after = Pt(0)
    r00 = p00.add_run("_______________________________\nAditya Godar\nStudent (PRN: 24070126014)\nClass: AIML A1\nSymbiosis Institute of Technology, Pune")
    r00.font.name = 'Times New Roman'
    r00.font.size = Pt(10)
    r00.font.color.rgb = COLOR_DARK_TEXT

    # Row 0, Cell 1: Project Mentor Signature Block
    c01 = sig_tbl.cell(0, 1)
    p01 = c01.paragraphs[0]
    p01.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p01.paragraph_format.line_spacing = 1.15
    p01.paragraph_format.space_before = Pt(0)
    p01.paragraph_format.space_after = Pt(0)
    r01 = p01.add_run("_______________________________\nProf. Darshan Talhande\nProject Mentor & Director,\nSEER Bharat Foundation, Pune\n[Official Seal & Signature]")
    r01.font.name = 'Times New Roman'
    r01.font.size = Pt(10)
    r01.font.color.rgb = COLOR_DARK_TEXT

    # Row 1, Cell 0: Faculty Coordinator Signature Block
    c10 = sig_tbl.cell(1, 0)
    p10 = c10.paragraphs[0]
    p10.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p10.paragraph_format.line_spacing = 1.15
    p10.paragraph_format.space_before = Pt(16)
    p10.paragraph_format.space_after = Pt(0)
    r10 = p10.add_run("_______________________________\nDr. Aditi Saxena\nFaculty Coordinator, Service Learning\nDepartment of AIML\nSymbiosis Institute of Technology, Pune")
    r10.font.name = 'Times New Roman'
    r10.font.size = Pt(10)
    r10.font.color.rgb = COLOR_DARK_TEXT

    # Row 1, Cell 1: Head of Department Signature Block
    c11 = sig_tbl.cell(1, 1)
    p11 = c11.paragraphs[0]
    p11.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p11.paragraph_format.line_spacing = 1.15
    p11.paragraph_format.space_before = Pt(16)
    p11.paragraph_format.space_after = Pt(0)
    r11 = p11.add_run("_______________________________\nDr. Sumit Kumar\nHead of Department, AIML\nDepartment of AIML\nSymbiosis Institute of Technology, Pune")
    r11.font.name = 'Times New Roman'
    r11.font.size = Pt(10)
    r11.font.color.rgb = COLOR_DARK_TEXT

    # Now move elements into place: cert_head_p -> cert_p1 -> cert_p2 -> cert_p3 -> cert_sp -> sig_tbl -> p_break
    curr = cert_head_p._element
    for elem in [cert_p1._element, cert_p2._element, cert_p3._element, cert_sp._element, sig_tbl._element]:
        curr.addnext(elem)
        curr = elem

    # Remove the empty placeholder paragraph
    p_empty._element.getparent().remove(p_empty._element)
    print("Certificate page content inserted cleanly!")

    # -------------------------------------------------------------
    # 2. FIX DUPLICATE 'Prof. Prof.' (5 locations)
    # -------------------------------------------------------------
    print("\n--- 2. Fixing duplicate 'Prof. Prof.' ---")
    fixed_count = 0
    for p_idx, p in enumerate(doc.paragraphs):
        if "Prof. Prof." in p.text:
            for r_idx in range(len(p.runs) - 1):
                if p.runs[r_idx].text.endswith("Prof. ") and p.runs[r_idx + 1].text.startswith("Prof. Darshan Talhande"):
                    p.runs[r_idx].text = p.runs[r_idx].text[:-6]  # remove trailing 'Prof. '
                    fixed_count += 1
                    print(f"Fixed 'Prof. Prof.' in doc.paragraphs[{p_idx}]: {repr(p.text[:80])}")
                    break
            # Fallback if in single run
            if "Prof. Prof." in p.text:
                for r in p.runs:
                    if "Prof. Prof." in r.text:
                        r.text = r.text.replace("Prof. Prof.", "Prof.")
                        fixed_count += 1
                        print(f"Fixed single-run 'Prof. Prof.' in doc.paragraphs[{p_idx}]: {repr(p.text[:80])}")

    # Check tables for Prof. Prof.
    for t_idx, tbl in enumerate(doc.tables):
        for r_idx, row in enumerate(tbl.rows):
            for c_idx, cell in enumerate(row.cells):
                if "Prof. Prof." in cell.text:
                    for p in cell.paragraphs:
                        for r_i in range(len(p.runs) - 1):
                            if p.runs[r_i].text.endswith("Prof. ") and p.runs[r_i + 1].text.startswith("Prof. Darshan Talhande"):
                                p.runs[r_i].text = p.runs[r_i].text[:-6]
                                fixed_count += 1
                                print(f"Fixed 'Prof. Prof.' in Table {t_idx} [R{r_idx}C{c_idx}]: {repr(cell.text[:80])}")
                                break
                        if "Prof. Prof." in cell.text:
                            for r in p.runs:
                                if "Prof. Prof." in r.text:
                                    r.text = r.text.replace("Prof. Prof.", "Prof.")
                                    fixed_count += 1
                                    print(f"Fixed single-run 'Prof. Prof.' in Table {t_idx} [R{r_idx}C{c_idx}]")

    print(f"Total 'Prof. Prof.' fixed: {fixed_count}")

    # -------------------------------------------------------------
    # 3. FIX LIST OF FIGURES FIGURE 34 (Preliminary Page 8)
    # -------------------------------------------------------------
    print("\n--- 3. Fixing List of Figures Figure 34 entry ---")
    for p_idx, p in enumerate(doc.paragraphs):
        if "Figure 34." in p.text and "Suvarna Patil" in p.text:
            for r in p.runs:
                if "Suvarna Patil" in r.text:
                    r.text = r.text.replace("Prof. Dr. Suvarna Patil", "Prof. Darshan Talhande")
                    r.text = r.text.replace("Suvarna Patil", "Darshan Talhande")
                    print(f"Updated doc.paragraphs[{p_idx}]: {repr(p.text)}")
            break

    # -------------------------------------------------------------
    # 4. FIX SECTION 1.2 PRONOUN & CAPITALIZATION MISMATCH
    # -------------------------------------------------------------
    print("\n--- 4. Fixing Section 1.2 Pronoun & Capitalization Mismatch ---")
    for p_idx, p in enumerate(doc.paragraphs):
        if "SEER Bharat Foundation is directed under the leadership of" in p.text and "Our mentor has spearheaded" in p.text:
            print(f"Found target paragraph at doc.paragraphs[{p_idx}]")
            # Update runs cleanly
            for r in p.runs:
                if "Prof. Dr. " in r.text:
                    r.text = r.text.replace("Prof. Dr. ", "Prof. ")
                if r.text == "Our":
                    r.text = "Prof. Talhande"
                if r.text == " mentor":
                    r.text = ""
                if "Under her guidance" in r.text:
                    r.text = r.text.replace("Under her guidance", "Under his guidance")
            print(f"Updated text: {repr(p.text[:140])}")
            break

    # -------------------------------------------------------------
    # 5. FIX SECTION 4.1 PROBLEM STATEMENT TYPO
    # -------------------------------------------------------------
    print("\n--- 5. Fixing Section 4.1 Problem Statement Typo ---")
    for p_idx, p in enumerate(doc.paragraphs):
        if "the Urban Ecosystem Health Intelligence System (UEHIS)" in p.text and "that ingests multi-spectral satellite imagery" in p.text:
            print(f"Found Problem Statement paragraph at doc.paragraphs[{p_idx}]")
            for r_i in range(len(p.runs)):
                if p.runs[r_i].text == ".":
                    p.runs[r_i].text = ","
                if p.runs[r_i].text.startswith(" that ingests"):
                    p.runs[r_i].text = " which ingests" + p.runs[r_i].text[len(" that ingests"):]
            print(f"Updated text: {repr(p.text[:140])}")
            break

    # -------------------------------------------------------------
    # 6. FIX SECTION 9 MISSING DESIGNATION & EN-DASHES
    # -------------------------------------------------------------
    print("\n--- 6. Fixing Section 9 Missing Designation & En-Dashes ---")
    for p_idx, p in enumerate(doc.paragraphs):
        if "The complete Annexure-3 record documents the full 15-week service engagement" in p.text:
            print(f"Found Section 9 paragraph at doc.paragraphs[{p_idx}]")
            for r in p.runs:
                if "\ufffd" in r.text:
                    r.text = r.text.replace("\ufffd", "–")
                if "and Faculty " in r.text and "Dr. Aditi" in p.text:
                    r.text = r.text.replace("and Faculty ", "and Faculty Coordinator ")
            print(f"Updated text: {repr(p.text[:140])}")
            break

    # -------------------------------------------------------------
    # 7. INSERT SECTION 2 SUBHEADINGS IN BODY
    # -------------------------------------------------------------
    print("\n--- 7. Inserting Section 2 Subheadings in Body ---")
    def create_h2(text):
        h = doc.add_paragraph()
        h.paragraph_format.keep_with_next = True
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        h.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = h.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = COLOR_NAVY
        return h

    # Subheading 2.1
    for p_idx, p in enumerate(doc.paragraphs):
        if p.text.startswith("Rapid demographic growth, economic expansion, and physical infrastructure densification"):
            # Check if 2.1 is already present right before it
            prev_p = doc.paragraphs[p_idx - 1] if p_idx > 0 else None
            if prev_p and "2.1 Urban Ecological Pressures" in prev_p.text:
                print("2.1 already present.")
            else:
                h21 = create_h2("2.1 Urban Ecological Pressures & Need for Spatial Intelligence")
                p._element.addprevious(h21._element)
                print(f"Inserted 2.1 subheading before doc.paragraphs[{p_idx}]")
            break

    # Subheading 2.2
    for p_idx, p in enumerate(doc.paragraphs):
        if p.text.startswith("The Urban Ecosystem Health Intelligence System (UEHIS) was conceived and developed as a collaborative"):
            prev_p = doc.paragraphs[p_idx - 1] if p_idx > 0 else None
            if prev_p and "2.2 Purpose and Civic Rationale" in prev_p.text:
                print("2.2 already present.")
            else:
                h22 = create_h2("2.2 Purpose and Civic Rationale of UEHIS")
                p._element.addprevious(h22._element)
                print(f"Inserted 2.2 subheading before doc.paragraphs[{p_idx}]")
            break

    # Subheading 2.3
    for p_idx, p in enumerate(doc.paragraphs):
        if p.text.startswith("Type of Work Done: Over the course of 120 officially documented"):
            prev_p = doc.paragraphs[p_idx - 1] if p_idx > 0 else None
            if prev_p and "2.3 Type of Work Done" in prev_p.text:
                print("2.3 already present.")
            else:
                h23 = create_h2("2.3 Type of Work Done & Project Synopsis")
                p._element.addprevious(h23._element)
                print(f"Inserted 2.3 subheading before doc.paragraphs[{p_idx}]")
            break

    # -------------------------------------------------------------
    # 8. SECTION 10 WORDING CLEANUP
    # -------------------------------------------------------------
    print("\n--- 8. Section 10 Wording Cleanup ---")
    for p_idx, p in enumerate(doc.paragraphs):
        if "Each image placeholder below represents a documented milestone" in p.text:
            print(f"Found Section 10 paragraph at doc.paragraphs[{p_idx}]")
            for r in p.runs:
                if "Each image placeholder below represents a documented milestone" in r.text:
                    r.text = r.text.replace(
                        "Each image placeholder below represents a documented milestone",
                        "Each photograph below documents an authentic milestone"
                    )
            print(f"Updated text: {repr(p.text[:140])}")
            break

    # -------------------------------------------------------------
    # SAVE UPDATED DOCUMENT
    # -------------------------------------------------------------
    print(f"\nSaving updated document to {DOCX_PATH}...")
    doc.save(DOCX_PATH)
    print("Saved successfully!")

if __name__ == "__main__":
    apply_all_updates()
