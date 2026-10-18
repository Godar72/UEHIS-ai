"""
Master Report Assembly Script
Builds 'UEHIS_Service_Learning_Report_Aditya_Godar.docx'
and converts it to 'UEHIS_Service_Learning_Report_Aditya_Godar.pdf' via Microsoft Word COM.
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from report_generator_core import COLOR_NAVY, COLOR_MUTED, COLOR_DARK_TEXT
from report_preliminaries import (
    create_document, build_cover_page, build_certificate_page,
    build_acknowledgement_page, build_abstract_page,
    build_table_of_contents, build_list_of_figures, build_list_of_tables
)
from report_part_a import build_part_a
from report_part_b import build_part_b
from report_annexures import build_annexures

def apply_headers_and_footers(doc):
    """Apply consistent academic header, footer, and page numbering fields."""
    for section in doc.sections:
        section.different_first_page_header_footer = True
        
        # Header (Pages 2+)
        header = section.header
        p_head = header.paragraphs[0]
        p_head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_head.paragraph_format.space_before = Pt(0)
        p_head.paragraph_format.space_after = Pt(0)
        r_h = p_head.add_run("UEHIS — Urban Ecosystem Health Intelligence System | Service Learning Report")
        r_h.font.name = 'Times New Roman'
        r_h.font.size = Pt(8.5)
        r_h.font.color.rgb = COLOR_MUTED
        
        # Footer (Pages 2+)
        footer = section.footer
        p_foot = footer.paragraphs[0]
        p_foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_foot.paragraph_format.space_before = Pt(0)
        p_foot.paragraph_format.space_after = Pt(0)
        
        # Text on left and page on right via spacing
        r_f1 = p_foot.add_run("Department of AIML, SIT Pune                                                                          Page ")
        r_f1.font.name = 'Times New Roman'
        r_f1.font.size = Pt(9)
        r_f1.font.color.rgb = COLOR_MUTED
        
        # Word Page Number Field
        fldSimple = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
        p_foot._p.append(fldSimple)

def main():
    print("=" * 60)
    print("Starting UEHIS Service Learning Report Generation...")
    print("=" * 60)

    doc = create_document()
    
    print("1. Assembling Preliminary Pages...")
    build_cover_page(doc)
    build_certificate_page(doc)
    build_acknowledgement_page(doc)
    build_abstract_page(doc)
    build_table_of_contents(doc)
    build_list_of_figures(doc)
    build_list_of_tables(doc)
    
    print("2. Assembling Part A: Service Learning Project Report...")
    build_part_a(doc)
    
    print("3. Assembling Part B: Technical Project Documentation...")
    build_part_b(doc)
    
    print("4. Assembling Appendices and Annexures...")
    build_annexures(doc)
    
    print("5. Applying Document Headers, Footers, and Page Fields...")
    apply_headers_and_footers(doc)
    
    docx_filename = "UEHIS_Service_Learning_Report_Aditya_Godar.docx"
    doc.save(docx_filename)
    print(f"Successfully saved Word document: {docx_filename}")
    print(f"File size: {os.path.getsize(docx_filename) / 1024:.1f} KB")

    print("\n6. Converting to PDF via Microsoft Word COM automation...")
    pdf_filename = "UEHIS_Service_Learning_Report_Aditya_Godar.pdf"
    
    try:
        import win32com.client
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        doc_path = os.path.abspath(docx_filename)
        pdf_path = os.path.abspath(pdf_filename)
        
        word_doc = word.Documents.Open(doc_path)
        # Update Table of Contents and fields if possible
        try:
            word_doc.Fields.Update()
        except Exception as e_field:
            print("Note on fields update:", e_field)
            
        # 17 = wdExportFormatPDF
        word_doc.ExportAsFixedFormat(pdf_path, 17)
        word_doc.Close(False)
        word.Quit()
        
        print(f"Successfully exported PDF: {pdf_filename}")
        print(f"PDF size: {os.path.getsize(pdf_path) / 1024:.1f} KB")
    except Exception as e:
        print(f"Word COM conversion encountered an issue: {e}")
        print("The .docx file is intact and ready.")

    print("=" * 60)
    print("UEHIS Report Generation Complete!")
    print("=" * 60)

if __name__ == "__main__":
    main()
