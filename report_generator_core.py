"""
UEHIS Service Learning Report Generator
Author: Antigravity Assistant for Aditya Godar (PRN: 24070126014)
Institution: Symbiosis Institute of Technology, Pune
Partner: SEER Bharat Foundation, Pune
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

# --- Color Constants ---
COLOR_NAVY = RGBColor(23, 43, 92)       # SEER Navy #172B5C
COLOR_GREEN = RGBColor(46, 155, 95)     # UEHIS Green #2E9B5F
COLOR_DARK_TEXT = RGBColor(30, 41, 59)  # Slate 800 #1E293B
COLOR_MUTED = RGBColor(100, 116, 139)   # Slate 500 #64748B
COLOR_LIGHT_BG = "F8FAFC"               # Slate 50
COLOR_BOX_BG = "F4F6F8"                 # Light Neutral Gray
COLOR_BORDER = "CBD5E1"                 # Slate 300
COLOR_HEADER_BG = "172B5C"              # SEER Navy Header
COLOR_ALT_ROW = "F1F5F9"                # Slate 100

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    """Set cell padding in twips."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    """Set background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}" w:val="clear"/>')
    tcPr.append(shd)

def set_cell_borders(cell, top="single", bottom="single", left="single", right="single", color="CBD5E1", sz="4"):
    """Set custom borders on a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    borders_xml = f'<w:tcBorders {nsdecls("w")}>'
    for side, val in [("top", top), ("left", left), ("bottom", bottom), ("right", right)]:
        if val == "none":
            borders_xml += f'<w:{side} w:val="none"/>'
        else:
            borders_xml += f'<w:{side} w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
    borders_xml += '</w:tcBorders>'
    tcPr.append(parse_xml(borders_xml))

def format_paragraph(p, space_before=0, space_after=6, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    """Apply consistent academic typography to a paragraph."""
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing

def add_styled_heading(doc, text, level):
    """Add professional heading with strict academic formatting."""
    h = doc.add_paragraph()
    h.paragraph_format.keep_with_next = True
    r = h.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.bold = True
    
    if level == 1:
        h.alignment = WD_ALIGN_PARAGRAPH.LEFT
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(6)
        r.font.size = Pt(15)
        r.font.color.rgb = COLOR_NAVY
    elif level == 2:
        h.alignment = WD_ALIGN_PARAGRAPH.LEFT
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        r.font.size = Pt(13)
        r.font.color.rgb = COLOR_NAVY
    elif level == 3:
        h.alignment = WD_ALIGN_PARAGRAPH.LEFT
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(3)
        r.font.size = Pt(11.5)
        r.font.color.rgb = COLOR_DARK_TEXT
    return h

def add_body_p(doc, text, bold_prefix=None, space_after=6):
    """Add standard academic body text."""
    p = doc.add_paragraph()
    format_paragraph(p, space_before=0, space_after=space_after, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(11)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_DARK_TEXT
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)
    r.font.color.rgb = COLOR_DARK_TEXT
    return p

def add_bullet_p(doc, text, bold_prefix=None):
    """Add academic bullet point."""
    p = doc.add_paragraph()
    format_paragraph(p, space_before=1, space_after=3, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    p.paragraph_format.left_indent = Inches(0.25)
    r_bullet = p.add_run("•  ")
    r_bullet.font.name = 'Times New Roman'
    r_bullet.font.bold = True
    r_bullet.font.color.rgb = COLOR_NAVY
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(11)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_DARK_TEXT
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)
    r.font.color.rgb = COLOR_DARK_TEXT
    return p

def add_figure_placeholder(doc, fig_num, title, dims, instructions, caption, image_path=None):
    """Add a full-width figure box, embedding real image if image_path exists, else placeholder."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    
    if image_path and os.path.exists(image_path):
        set_cell_borders(cell, top="single", bottom="single", left="single", right="single", color="CBD5E1", sz="4")
        set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run()
        run.add_picture(image_path, width=Inches(6.2))
    else:
        set_cell_shading(cell, COLOR_BOX_BG)
        set_cell_borders(cell, top="single", bottom="single", left="single", right="single", color="94A3B8", sz="6")
        set_cell_margins(cell, top=160, bottom=160, left=200, right=200)
        
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        r1 = p.add_run(f"[INSERT FIGURE {fig_num} HERE — {title.upper()}]")
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(10.5)
        r1.font.bold = True
        r1.font.color.rgb = COLOR_NAVY
        
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_before = Pt(2)
        p2.paragraph_format.space_after = Pt(3)
        r2 = p2.add_run(f"Suggested Dimensions: {dims}")
        r2.font.name = 'Times New Roman'
        r2.font.size = Pt(9)
        r2.font.color.rgb = COLOR_MUTED
        
        p3 = cell.add_paragraph()
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p3.paragraph_format.space_before = Pt(2)
        p3.paragraph_format.space_after = Pt(8)
        r3 = p3.add_run(f"Evidence to Insert: {instructions}")
        r3.font.name = 'Times New Roman'
        r3.font.size = Pt(9)
        r3.font.italic = True
        r3.font.color.rgb = COLOR_DARK_TEXT
    
    # Caption below box
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(4)
    p_cap.paragraph_format.space_after = Pt(10)
    p_cap.paragraph_format.keep_with_next = False
    r_cap = p_cap.add_run(f"Figure {fig_num}. {caption}")
    r_cap.font.name = 'Times New Roman'
    r_cap.font.size = Pt(9.5)
    r_cap.font.italic = True
    r_cap.font.color.rgb = COLOR_DARK_TEXT

def add_side_by_side_placeholders(doc, fig1, fig2):
    """Add two clean side-by-side boxes, embedding real images if image_path exists, else placeholders."""
    tbl = doc.add_table(rows=2, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    col_widths = [Inches(3.1), Inches(0.3), Inches(3.1)]
    for row in tbl.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

    # Row 0: Images or Placeholders
    for col_idx, fig in [(0, fig1), (2, fig2)]:
        cell = tbl.cell(0, col_idx)
        img_path = fig.get('image_path')
        if img_path and os.path.exists(img_path):
            set_cell_borders(cell, top="single", bottom="single", left="single", right="single", color="CBD5E1", sz="4")
            set_cell_margins(cell, top=60, bottom=60, left=60, right=60)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run()
            run.add_picture(img_path, width=Inches(3.0))
        else:
            set_cell_shading(cell, COLOR_BOX_BG)
            set_cell_borders(cell, top="single", bottom="single", left="single", right="single", color="94A3B8", sz="6")
            set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
            
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(f"[INSERT FIG. {fig['num']} — {fig['title'].upper()}]")
            r.font.name = 'Times New Roman'
            r.font.size = Pt(9.5)
            r.font.bold = True
            r.font.color.rgb = COLOR_NAVY
            
            p2 = cell.add_paragraph()
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p2.paragraph_format.space_before = Pt(2)
            p2.paragraph_format.space_after = Pt(2)
            r2 = p2.add_run(f"Dims: {fig['dims']}")
            r2.font.name = 'Times New Roman'
            r2.font.size = Pt(8.5)
            r2.font.color.rgb = COLOR_MUTED
            
            p3 = cell.add_paragraph()
            p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p3.paragraph_format.space_before = Pt(2)
            p3.paragraph_format.space_after = Pt(6)
            r3 = p3.add_run(f"Evidence: {fig['evidence']}")
            r3.font.name = 'Times New Roman'
            r3.font.size = Pt(8.5)
            r3.font.italic = True
            r3.font.color.rgb = COLOR_DARK_TEXT

    # Spacer cell styling in Row 0 and Row 1 (clear borders)
    for r_idx in [0, 1]:
        spacer_cell = tbl.cell(r_idx, 1)
        set_cell_borders(spacer_cell, top="none", left="none", bottom="none", right="none")

    # Row 1: Captions
    for col_idx, fig in [(0, fig1), (2, fig2)]:
        cell = tbl.cell(1, col_idx)
        set_cell_borders(cell, top="none", left="none", bottom="none", right="none")
        set_cell_margins(cell, top=60, bottom=80, left=40, right=40)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(8)
        r = p.add_run(f"Figure {fig['num']}. {fig['caption']}")
        r.font.name = 'Times New Roman'
        r.font.size = Pt(9)
        r.font.italic = True
        r.font.color.rgb = COLOR_DARK_TEXT

def create_styled_table(doc, headers, data, col_widths, table_num, caption):
    """Create a fully styled academic table with Navy headers and clear cell padding."""
    # Table title above table
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(10)
    p_title.paragraph_format.space_after = Pt(4)
    p_title.paragraph_format.keep_with_next = True
    r_t = p_title.add_run(f"Table {table_num}. {caption}")
    r_t.font.name = 'Times New Roman'
    r_t.font.size = Pt(10)
    r_t.font.bold = True
    r_t.font.color.rgb = COLOR_NAVY
    
    tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    # Header Row
    hdr_row = tbl.rows[0]
    trPr = hdr_row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    
    for c_idx, h_text in enumerate(headers):
        cell = hdr_row.cells[c_idx]
        cell.width = col_widths[c_idx]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, COLOR_HEADER_BG)
        set_cell_borders(cell, top="single", bottom="single", left="single", right="single", color="CBD5E1", sz="4")
        set_cell_margins(cell, top=140, bottom=140, left=140, right=140)
        
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h_text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    # Data Rows
    for r_idx, row_values in enumerate(data):
        row = tbl.rows[r_idx + 1]
        bg_col = COLOR_ALT_ROW if (r_idx % 2 == 1) else "FFFFFF"
        for c_idx, val in enumerate(row_values):
            cell = row.cells[c_idx]
            cell.width = col_widths[c_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_shading(cell, bg_col)
            set_cell_borders(cell, top="single", bottom="single", left="single", right="single", color="E2E8F0", sz="4")
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(str(val))
            r.font.name = 'Times New Roman'
            r.font.size = Pt(9)
            r.font.color.rgb = COLOR_DARK_TEXT

    # Trailing spacing after table
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(6)

print("Report generator core utilities defined successfully.")
